"""Saves extracted data to an Excel file."""
import shutil

import openpyxl
from openpyxl.styles import NamedStyle

def find_table(workbook, table_name):
    """Identifies the worksheet containing the desired table."""
    for worksheet in workbook.worksheets:
        for table in worksheet.tables.values():
            if table.name == table_name:
                return table, worksheet

def save_data(data, config, log):
    """Saves extracted data to Excel file."""
    log.info(f'Creating backup Excel file: {config['excel_backup_path']}')
    shutil.copy(config['excel_file_path'], config['excel_backup_path'])

    log.info(f'Opening Excel file to save data: {config['excel_file_path']}')
    workbook = openpyxl.load_workbook(config['excel_file_path'])

    # Collection of styles
    number_format = '0.00'
    currency_format = '$0.00'
    date_format = 'YYYY-MMM-DD'

    # Dictionary outlining extracted data to table names
    excel_key_dict = {
        'paycheque_details': 'paycheque_details',
        'baseline_details': 'baseline_details',
        'tax_data': 'tax_data',
        'hours_and_earnings': 'hours_and_earnings',
        'taxes': 'taxes',
        'before_tax_deductions': 'before_tax_deductions',
        'after_tax_deductions': 'after_tax_deductions',
        'employer_paid_benefits': 'employer_paid_benefits',
        'gross_and_net': 'gross_and_net_pay',
        'vacation': 'vacation',
        'bank_balances': 'bank_balances',
        'advance_outstanding': 'advance_outstanding',
        'direct_deposit_distribution': 'direct_deposit_distribution',
        'net_pay_distribution': 'net_pay_distribution',
        'message': 'message',
    }

    # Collect the paycheque details; these are applied to every table
    paycheque_details = data['paycheque_details'][0]

    for dict_key, table_name in excel_key_dict.items():
        log.info(f'  Adding data to table: {table_name}')

        # Find required table
        table, worksheet = find_table(workbook, table_name)

        # Calculate the next row
        last_row = table.ref.split(":")[1][1:]  # Get the last row number, e.g., "5" from "A1:D5"
        next_row = int(last_row) + 1

        # Append data row in the table
        data_values = data[dict_key]

        for row, item_data in enumerate(data_values, start=next_row):
            log.debug('Adding data to row {row}')

            # Add paycheque details to all other tables
            if dict_key != 'paycheque_details':
                item_data = paycheque_details + item_data

            for col, cell_data in enumerate(item_data, start=1):
                cell = worksheet.cell(row=row, column=col)
                cell.value = cell_data['value']

                if cell_data['data_type'] == 'number':
                    cell.number_format = number_format
                elif cell_data['data_type'] == 'currency':
                    cell.number_format = currency_format
                if cell_data['data_type'] == 'date':
                    cell.number_format = date_format

        # Update table dimensions to include new row
        top_left, bottom_right = table.ref.split(':')
        bottom_right_col, _ = openpyxl.utils.cell.coordinate_from_string(bottom_right)
        new_bottom_right = f'{bottom_right_col}{row}'
        new_range = f'{top_left}:{new_bottom_right}'

        log.debug(f'Extending table range to {new_bottom_right}')
        worksheet.tables[table_name].ref = new_range

    # Save the workbook
    workbook.save(config['excel_file_path'])
